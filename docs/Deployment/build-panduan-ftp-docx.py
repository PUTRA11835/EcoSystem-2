# -*- coding: utf-8 -*-
"""
Membangkitkan: PANDUAN-FTP-DAN-DEPLOY-PEMULA.docx (arahan untuk pemilik, 22 Okt 2026).
Aturan: .docx = arahan untuk pemilik; ubah data di skrip ini lalu jalankan ulang (timpa satu berkas):
    py docs/Deployment/build-panduan-ftp-docx.py
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "PANDUAN-FTP-DAN-DEPLOY-PEMULA.docx")
HEAD = "1F3A5F"


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def table(doc, header, rows, widths):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(text)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, HEAD)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(str(v)).font.size = Pt(9.5)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph()


def para(doc, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    return p


def steps(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Number")


def bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def code(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)


doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)

doc.add_heading("Panduan FTP dan Deploy untuk Pemula", 0)
para(doc, "EcoSystem-2 · cPanel ecosystemtest.org · disusun 22 Oktober 2026. Dokumen perencanaan: belum ada kode aplikasi atau pengaturan server yang diubah.")

doc.add_heading("1. Apa yang terlihat dari tangkapan layar Anda", 1)
table(doc, ["Temuan", "Artinya"], [
    ["Aplikasi ada di /home/ecos4415/EcoSystem-2 (di luar public_html)", "Bagus: berkas penting (.env, app, vendor) tidak berada di folder web. Yang perlu dipastikan: folder web (document root) domain Anda menunjuk ke mana (lihat bagian 3)."],
    ["Ada folder .git, .claude, docs, tests di server", "Seluruh folder kerja di-zip lalu di-upload. Folder .git di server berisi catatan commit yang diunggah, jadi bisa dipakai untuk mengetahui versi yang terpasang (bagian 4)."],
    ["Izin folder 0777 dan berkas 0666 (termasuk .env)", "Terlalu longgar. Di hosting bersama, orang lain di server yang sama berpotensi bisa membaca .env (berisi sandi database dan kunci aplikasi). Dianjurkan folder 755, berkas 644, .env 600 (bagian 5)."],
    ["Ada EcoSystem-2-main, EcoSystem-2-main1, JARVIES-new-main", "Salinan lama. Jangan dihapus dulu; catat saja. Pastikan tidak ada yang bisa dibuka dari web."],
    ["vendor ada di server", "Dependensi sudah terpasang. Selama composer.json/composer.lock tidak berubah, deploy cukup mengunggah kode, tanpa vendor."],
], [6.5, 10])

doc.add_heading("2. Cek keamanan 2 menit (lakukan sekarang)", 1)
steps(doc, [
    "Buka di peramban: https://ALAMAT-APLIKASI-ANDA/.env (ganti dengan alamat yang biasa Anda pakai untuk membuka aplikasi).",
    "Lalu buka juga https://ALAMAT-APLIKASI-ANDA/composer.json dan https://ALAMAT-APLIKASI-ANDA/storage/logs/laravel.log.",
    "Hasil yang BENAR: halaman 403, 404, atau halaman galat aplikasi. Hasil yang SALAH: isi berkas tampil di layar. Bila salah, hentikan semuanya dan kabari saya sekarang, karena berarti sandi database terbuka untuk umum.",
])

doc.add_heading("3. Bagian A — Membuat akun FTP khusus deploy", 1)
para(doc, "Mengapa akun khusus: akun utama (ecos4415) memegang seluruh hosting. Bila sandinya bocor, semuanya terbuka. Akun khusus hanya bisa masuk ke satu folder dan bisa dihapus kapan saja tanpa mengganggu yang lain.")
steps(doc, [
    "cPanel → FTP Accounts (halaman yang Anda tangkap tadi).",
    "Log In: ketik deploy (huruf kecil, tanpa spasi). Kotak merah \"You must enter an FTP username\" akan hilang setelah diisi.",
    "Domain: biarkan domain yang tampil, atau pilih domain utama. Nama lengkap login akan menjadi deploy@NAMA-DOMAIN — catat persis seperti yang tampil.",
    "Password: klik Password Generator, centang \"I have copied this password in a safe place\", lalu Use Password. Simpan di pengelola sandi Anda. JANGAN tempel sandi ini ke chat.",
    "Directory: kotak ini sudah berawalan /home/ecos4415/. Hapus isi default (public_html/hrm.ecosystemtest.org/) lalu ketik: EcoSystem-2 . Artinya akun ini hanya bisa melihat folder aplikasi.",
    "Quota: pilih Unlimited (aplikasi sudah memakai lebih dari 1,7 GB).",
    "Klik Create FTP Account. Akun muncul di tabel FTP Accounts.",
    "Di baris akun baru klik Configure FTP Client. Catat: FTP Server (host) dan FTP & Explicit FTPS Port (biasanya 21). Bagian \"FTP over TLS\" biasanya tertulis tersedia. Foto bagian ini boleh dikirim ke saya (tutupi sandi).",
])
para(doc, "Catatan folder: jika folder web domain Anda ternyata terpisah dari EcoSystem-2 (mis. isi public/ disalin ke public_html/...), deploy butuh dua tujuan. Saya akan memastikannya setelah Anda mengirim jawaban bagian 7.")

doc.add_heading("4. Bagian B — Mengetahui versi yang terpasang di server", 1)
para(doc, "Supaya saya bisa menghitung berkas apa saja yang berbeda antara server dan kode terbaru:")
steps(doc, [
    "File Manager → buka EcoSystem-2 → aktifkan Settings (kanan atas) → Show Hidden Files (dotfiles).",
    "Buka folder .git → klik berkas HEAD → View. Salin isinya (satu baris, mis. ref: refs/heads/aldy_hr).",
    "Masih di .git, buka folder logs → berkas HEAD → View. Salin BARIS TERAKHIR (berawalan dua kode panjang, memuat tanggal dan pesan commit).",
    "Kirim kedua teks itu ke saya. Sandi tidak ada di sana.",
])

doc.add_heading("5. Bagian C — Merapikan izin berkas (disarankan, aman)", 1)
table(doc, ["Apa", "Izin sekarang", "Izin yang dianjurkan", "Cara"], [
    [".env", "0666", "0600", "File Manager → klik kanan .env → Permissions → centang Read dan Write hanya di baris User → Change Permissions"],
    ["storage dan bootstrap/cache", "0777", "0775", "Sama; ubah satu per satu"],
    ["folder lain", "0777", "0755", "Opsional; lebih mudah lewat cron sekali jalan (saya siapkan perintahnya bila Anda mau)"],
], [3.5, 3, 3.5, 6.5])
para(doc, "Setelah mengubah .env, buka satu halaman aplikasi. Jika muncul galat 500, kembalikan .env ke 0644 dan kabari saya. Di cPanel PHP berjalan sebagai akun Anda sendiri, jadi 0600 seharusnya tetap terbaca oleh aplikasi.")

doc.add_heading("6. Bagian D — Mencoba FTPS dari komputer Anda", 1)
steps(doc, [
    "Unduh FileZilla Client (filezilla-project.org), pasang.",
    "File → Site Manager → New Site. Protocol: FTP. Host: yang tercatat di Configure FTP Client. Port: 21. Encryption: Require explicit FTP over TLS. Logon Type: Normal. User: deploy@NAMA-DOMAIN. Password: sandi tadi.",
    "Klik Connect. Bila muncul jendela \"Unknown certificate\", centang Always trust lalu OK (catat: nama pada sertifikat harus sama dengan host; bila berbeda, kabari saya, karena GitHub tidak akan ikut percaya).",
    "Berhasil bila panel kanan menampilkan isi EcoSystem-2 (app, config, public, dst.). Jangan unggah/hapus apa pun saat mencoba.",
    "Cek cepat port dari PowerShell: Test-NetConnection NAMA-HOST -Port 21 → TcpTestSucceeded harus True.",
])
para(doc, "Jika langkah 3 menolak TLS atau port tertutup, FTPS tidak tersedia di paket hosting Anda. Itu bukan akhir: lanjut ke bagian 8.")

doc.add_heading("7. Yang saya butuhkan dari Anda", 1)
table(doc, ["#", "Info", "Cara mendapatkan"], [
    ["1", "Hasil cek keamanan bagian 2 (tiga alamat)", "Peramban"],
    ["2", "Isi .git/HEAD dan baris terakhir .git/logs/HEAD", "Bagian 4"],
    ["3", "Document root domain aplikasi", "cPanel → Domains → baris domain → kolom Document Root (kirim teksnya)"],
    ["4", "Host dan port FTP, status FTPS", "Bagian 3 langkah 8 dan bagian 6"],
    ["5", "Apakah menu Git Version Control ada, dan apakah ada repositori di dalamnya", "cPanel → Files → Git Version Control"],
    ["6", "Apakah Anda punya hak Admin di repositori GitHub PUTRA11835/EcoSystem-2", "GitHub → repositori → tab Settings muncul atau tidak. Rahasia (Secrets) hanya bisa dibuat oleh admin; bila tidak ada, minta pemilik repo"],
], [1, 7, 8.5])

doc.add_heading("8. Alur akhir yang akan kita bangun (setelah bagian 7 terjawab)", 1)
table(doc, ["Tahap", "Isi", "Siapa"], [
    ["1", "Simpan host, user, dan sandi FTP sebagai GitHub Secrets (bukan di kode)", "Anda (admin repo)"],
    ["2", "Saya buat berkas .github/workflows/deploy.yml: tombol manual \"Run workflow\", mengunggah hanya berkas yang berubah, mengecualikan .env, storage, vendor, .git, docs, tests, .claude. Percobaan pertama dry-run (tidak mengunggah apa pun)", "Saya"],
    ["3", "Percobaan ke satu berkas kecil, lalu deploy sungguhan", "Anda menekan tombol, saya baca hasilnya"],
    ["4", "Migrasi database: cron cPanel sekali jalan (php artisan migrate --force), dihapus setelah selesai; cadangkan DB lebih dulu", "Anda, dengan perintah dari saya"],
], [1.5, 11, 4])
para(doc, "Bila FTPS tidak tersedia atau diblokir: pakai cPanel Git Version Control (Pull dari GitHub lalu Deploy HEAD Commit) atau tetap zip berkas yang berubah. Saya bisa membuat zip berisi hanya berkas berbeda dengan struktur folder yang benar, tinggal Extract di File Manager.")

doc.add_heading("9. Aturan aman yang tetap berlaku", 1)
bullets(doc, [
    "Jangan pernah mengirim sandi FTP/database ke chat atau menyimpannya di repositori.",
    "Cadangkan database (phpMyAdmin → Export) sebelum SQL atau migrasi apa pun.",
    "Deploy hanya setelah diuji di lokal; migrasi selalu terpisah dan terjadwal.",
    "Setiap perubahan langsung di server dicatat di docs/HRGA/09-PROTOKOL-PERBAIKAN-UAT.md bagian 5 dan dicerminkan ke repositori.",
])

doc.save(OUT)
print("ditulis:", OUT)
